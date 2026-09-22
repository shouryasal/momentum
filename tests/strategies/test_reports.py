"""walk_forward window math + zip parsing, g2 evaluation, benchmark rows, excel view."""

import json
import zipfile
from datetime import date

import pandas as pd
import pytest

from runs import g2_check, sleeve_c_benchmark, walk_forward


class TestWalkForward:
    def test_windows_expanding(self):
        w = walk_forward.windows("20210101", 6, date(2024, 1, 15))
        assert w[0] == ("20220101", "20220701")
        assert w[-1][1] == "20240115"
        # contiguous
        for (_s1, e1), (s2, _e2) in zip(w, w[1:], strict=False):
            assert e1 == s2

    def test_parse_backtest_zip(self, tmp_path):
        payload = {
            "strategy": {"SleeveA": {
                "profit_total": 0.42, "max_drawdown_account": 0.18,
                "total_trades": 37, "backtest_start": "2024-01-01 00:00:00",
                "backtest_end": "2024-07-01 00:00:00",
            }}
        }
        z = tmp_path / "backtest-result-2026-09-22.zip"
        with zipfile.ZipFile(z, "w") as zf:
            zf.writestr("backtest-result-2026-09-22.json", json.dumps(payload))
            zf.writestr("backtest-result-2026-09-22_config.json", "{}")
        got = walk_forward.parse_backtest_zip(z)
        assert got["profit_total_pct"] == pytest.approx(42.0)
        assert got["max_drawdown_pct"] == pytest.approx(18.0)
        assert got["trades"] == 37


def _candles(start, days, daily_ret):
    dates = pd.date_range(start, periods=days, freq="1D", tz="UTC")
    closes = 100 * (1 + daily_ret) ** pd.RangeIndex(days).to_series()
    return pd.DataFrame({"date": dates, "close": closes.values,
                         "open": closes.values, "high": closes.values,
                         "low": closes.values, "volume": 1.0})


class TestG2:
    def test_btc_hold_metrics(self):
        c = _candles("2024-01-01", 200, 0.01)
        m = g2_check.btc_hold_metrics(c, "2024-01-01", "2024-06-01")
        assert m["return_pct"] > 0 and m["max_drawdown_pct"] == pytest.approx(0.0)

    def test_evaluate_thresholds(self):
        btc = {"return_pct": 100.0, "max_drawdown_pct": 30.0}
        good = {"profit_total_pct": 70.0, "max_drawdown_pct": 20.0}
        v = g2_check.evaluate(good, btc, cost_pct_month=0.5)
        assert v["pass"]
        v = g2_check.evaluate({"profit_total_pct": 50.0, "max_drawdown_pct": 20.0}, btc, 0.5)
        assert not v["pass"] and not v["checks"]["return_ratio_ok"]
        v = g2_check.evaluate({"profit_total_pct": 70.0, "max_drawdown_pct": 35.0}, btc, 0.5)
        assert not v["checks"]["dd_better_than_btc"]


class TestBenchmark:
    def test_rows_math(self):
        c = _candles("2026-10-01", 60, 0.01)
        rows = sleeve_c_benchmark.benchmark_rows(c, "2026-10-27", 10000, 15, "2026-11-05")
        assert rows[0][0] == "2026-10-27"
        assert rows[0][1] == pytest.approx(10000 * (1 - 0.0015))
        # NAV tracks close ratio
        assert rows[-1][1] / rows[0][1] == pytest.approx(1.01 ** (len(rows) - 1))

    def test_empty_before_start(self):
        c = _candles("2026-01-01", 10, 0.0)
        assert sleeve_c_benchmark.benchmark_rows(c, "2026-10-27", 10000, 15) == []


def test_excel_view_writes_all_sheets(tmp_path):
    from openpyxl import load_workbook

    from ops import db
    from ops.config import load_config
    from runs.excel_view import write_workbook

    cfg = load_config()
    journal, _ = db.init_all(cfg, root=tmp_path)
    with db.connect(journal) as c:
        c.execute("INSERT INTO nav_daily(date_utc, sleeve, nav_usdt) VALUES"
                  " ('2026-10-27','a',10000), ('2026-10-27','b',10000),"
                  " ('2026-10-27','benchmark',9985)")
        c.execute("INSERT INTO fills(ts_utc, sleeve, pair, side, fill_amount, fill_price,"
                  " quote_bid, quote_ask) VALUES"
                  " ('2026-10-27T04:30:00Z','a','BTC/USDT','buy',0.01,100.1,100.0,100.2)")
        c.execute("INSERT INTO proposals(run_id, shadow, ts_utc, valid, module,"
                  " targets_json, abstain) VALUES"
                  " ('2026-10-27T08:30+04:00',0,'2026-10-27T04:30:00Z',1,'trend',"
                  "'{\"BTC\":0.4}',0)")
        c.execute("INSERT INTO runs(run_id, stage, kind, started_utc, served_model,"
                  " effort, escalated, trigger_reason, status) VALUES"
                  " ('2026-10-27T08:30+04:00','decide','research','2026-10-27T04:30:00Z',"
                  "'claude-opus-5','max',0,'news:hack','success')")
        c.execute("INSERT INTO decision_grades(run_id, graded_at, review_week,"
                  " process_grade, process_rubric_json, outcome_grade, grader_model,"
                  " grader_run_id) VALUES ('2026-10-27T08:30+04:00','x','2026-W44',55,"
                  "'{}','worse','m','r')")
        c.execute("INSERT INTO whatif_nav(date_utc, nav_usdt, turnover, cost_usdt)"
                  " VALUES ('2026-10-27', 10120.5, 0.65, 9.75)")
        c.commit()

    class P:  # point paths at the tmp journal
        pass

    cfg2 = cfg.model_copy(deep=True)
    cfg2.paths.journal_db = str(journal.relative_to(tmp_path))
    import runs.excel_view as ev
    orig = ev.REPO_ROOT
    ev.REPO_ROOT = tmp_path
    try:
        out = write_workbook(cfg2, out=tmp_path / "earn.xlsx")
    finally:
        ev.REPO_ROOT = orig
    wb = load_workbook(out)
    assert {"NAV", "Trades", "Costs", "Gate", "Limits",
            "Decisions", "Mistakes", "WhatIf"} <= set(wb.sheetnames)
    trades = wb["Trades"]
    # slippage bps for the buy: (100.1 - 100.1)/100.1 -> 0
    assert trades.cell(row=2, column=9).value == pytest.approx(0.0, abs=1e-6)
    dec = wb["Decisions"]
    row2 = [dec.cell(row=2, column=i).value for i in range(1, 17)]
    assert row2[0] == "2026-10-27T08:30+04:00" and row2[2] == "triggered"
    assert row2[3] == "claude-opus-5" and row2[4] == "max"
    assert row2[13] == 55 and "process 55" in row2[14]      # wrong: low grade
    mis = wb["Mistakes"]
    assert mis.cell(row=2, column=1).value == "2026-10-27T08:30+04:00"
    assert "process 55" in mis.cell(row=2, column=2).value
    wi = wb["WhatIf"]
    assert wi.cell(row=2, column=1).value == "2026-10-27"
    assert wi.cell(row=2, column=2).value == pytest.approx(10120.5)
    assert wi.cell(row=2, column=3).value == pytest.approx(10000)   # sleeve b nav
